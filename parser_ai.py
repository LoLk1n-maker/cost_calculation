"""
parser_ai.py — БЛОК A1 «Анализ входных данных».

Задача блока: принять исходные документы (спецификация + технологический маршрут)
и привести их к структурированным данным (списки Material и Operation).

В MVP реализовано чтение из структурированного Excel. Метод parse_with_llm() —
точка расширения: сюда подключается разбор «сырых» файлов через LLM (см. README).
"""
import openpyxl
from models import Material, Operation, ProductData


def _cell(row, idx, default=""):
    v = row[idx].value if idx < len(row) else None
    return v if v is not None else default


def parse_specification(path: str):
    """Читает спецификацию (BOM) из Excel → список Material.
    Ожидаемые колонки: Наименование | Марка | Масса заготовки, кг | Масса детали, кг
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    materials = []
    for row in ws.iter_rows(min_row=2):  # строка 1 — заголовок
        name = _cell(row, 0)
        if not name:
            continue
        materials.append(Material(
            name=str(name),
            grade=str(_cell(row, 1)),
            mass_billet=float(_cell(row, 2, 0) or 0),
            mass_part=float(_cell(row, 3, 0) or 0),
        ))
    return materials


def parse_route(path: str):
    """Читает технологический маршрут из Excel → список Operation.
    Ожидаемые колонки: Операция | Оборудование | Норма времени, ч | Разряд рабочего
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    operations = []
    for row in ws.iter_rows(min_row=2):
        name = _cell(row, 0)
        if not name:
            continue
        operations.append(Operation(
            name=str(name),
            equipment=str(_cell(row, 1)),
            time_norm=float(_cell(row, 2, 0) or 0),
            worker_grade=str(_cell(row, 3)),
        ))
    return operations


def analyze_inputs(spec_path: str, route_path: str, product_name: str) -> ProductData:
    """Основная функция блока A1: собирает структурированные данные изделия."""
    materials = parse_specification(spec_path)
    operations = parse_route(route_path)
    return ProductData(product_name=product_name,
                       materials=materials,
                       operations=operations)


def parse_with_llm(raw_text: str) -> dict:
    """
    ТОЧКА РАСШИРЕНИЯ (AI-парсер).
    Для «сырых» документов (PDF/произвольный Excel) здесь вызывается LLM,
    который возвращает структурированный JSON с материалами и операциями.

    Пример интеграции с Anthropic API:

        import anthropic, json
        client = anthropic.Anthropic(api_key="...")
        msg = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=2000,
            messages=[{"role": "user", "content":
                "Извлеки материалы и операции из текста и верни строго JSON "
                "{materials:[...], operations:[...]}. Текст:\\n" + raw_text}],
        )
        return json.loads(msg.content[0].text)

    В текущем MVP не используется — данные читаются из структурированного Excel.
    """
    raise NotImplementedError("LLM-парсинг подключается по инструкции в README.")
